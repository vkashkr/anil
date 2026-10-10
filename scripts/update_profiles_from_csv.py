import csv
import json
from datetime import datetime
from pathlib import Path

import boto3
from botocore.exceptions import ClientError
from openpyxl import load_workbook

TABLE_NAME = 'gif-gif'
BUCKET_NAME = 'gif-gif'
CSV_PATH = r"C:\Users\Vikash\OneDrive\Documents\profiles.xlsx"


def convert_dynamodb_like(value):
    if value is None:
        return None
    if isinstance(value, dict):
        if set(value.keys()) == {'S'}:
            return value['S']
        if set(value.keys()) == {'N'}:
            return value['N']
        if set(value.keys()) == {'BOOL'}:
            return value['BOOL']
        if set(value.keys()) == {'L'}:
            return [convert_dynamodb_like(v) for v in value['L']]
        if set(value.keys()) == {'M'}:
            return {k: convert_dynamodb_like(v) for k, v in value['M'].items()}
        if set(value.keys()) == {'SS'}:
            return value['SS']
        return {k: convert_dynamodb_like(v) for k, v in value.items()}
    if isinstance(value, list):
        return [convert_dynamodb_like(v) for v in value]
    return value


def parse_jsonish(raw):
    if raw is None:
        return None
    text = str(raw).strip()
    if text == '':
        return None
    lowered = text.lower()
    if lowered == 'true':
        return True
    if lowered == 'false':
        return False
    if text.startswith('{') or text.startswith('['):
        try:
            return convert_dynamodb_like(json.loads(text))
        except Exception:
            return text
    return text


def item_from_row(row):
    profile_id = str(row.get('id') or row.get('PK') or '').strip()
    if not profile_id:
        return None

    item = {
        'PK': str(profile_id),
        'updatedAt': row.get('updatedAt') or datetime.utcnow().isoformat(),
    }

    for key in [
        'age', 'city', 'country', 'customCss', 'description', 'district', 'gender',
        'location', 'name', 'place', 'seoDescription', 'seoTitle', 'services', 'state'
    ]:
        raw = row.get(key)
        if raw is None:
            continue
        value = parse_jsonish(raw)
        if value is None or value == '':
            continue
        if key == 'age':
            try:
                item[key] = int(str(value))
            except ValueError:
                item[key] = str(value)
        elif key == 'services':
            service_values = value if isinstance(value, list) else str(value).split(',')
            item[key] = [
                str(service).strip()
                for service in service_values
                if service is not None and str(service).strip()
            ]
        else:
            item[key] = value

    visible_raw = row.get('isVisible')
    if visible_raw is not None:
        item['isVisible'] = str(visible_raw).strip().lower() == 'true'

    for key in ['extraProperties', 'images', 'reviews']:
        raw = row.get(key)
        if raw is None or str(raw).strip() == '':
            continue
        parsed = parse_jsonish(raw)
        if isinstance(parsed, (list, dict)):
            item[key] = parsed

    return item


def read_rows(path):
    if Path(path).suffix.lower() == '.xlsx':
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            worksheet_rows = workbook.active.iter_rows(values_only=True)
            headers = next(worksheet_rows, None)
            if headers is None:
                return

            fieldnames = [
                str(header).strip() if header is not None else ''
                for header in headers
            ]
            for values in worksheet_rows:
                yield {
                    fieldname: value
                    for fieldname, value in zip(fieldnames, values)
                    if fieldname
                }
        finally:
            workbook.close()
        return

    with open(path, newline='', encoding='utf-8-sig') as handle:
        yield from csv.DictReader(handle)


def ascii_safe_metadata(value):
    if value is None:
        return ''
    text = str(value)
    text = text.replace('\u2013', '-').replace('\u2014', '-').replace('\u2018', "'").replace('\u2019', "'").replace('\u201c', '"').replace('\u201d', '"')
    return text.encode('ascii', 'ignore').decode('ascii').strip()


def update_s3_metadata(s3_client, profile_id, metadata):
    candidates = [
        f'{profile_id}/profile.jpg',
    ]
    for key in candidates:
        try:
            head = s3_client.head_object(Bucket=BUCKET_NAME, Key=key)
            s3_client.copy_object(
                Bucket=BUCKET_NAME,
                Key=key,
                CopySource={'Bucket': BUCKET_NAME, 'Key': key},
                Metadata=metadata,
                MetadataDirective='REPLACE',
                ContentType=head.get('ContentType'),
            )
            return key
        except ClientError:
            continue
    return None


def main():
    dynamodb = boto3.resource('dynamodb')
    s3 = boto3.client('s3', region_name='us-east-1')
    table = dynamodb.Table(TABLE_NAME)

    rows_updated = 0
    s3_updated = 0

    for row in read_rows(CSV_PATH):
        item = item_from_row(row)
        if not item:
            continue

        table.put_item(Item=item)
        rows_updated += 1

        meta = {}
        if item.get('name'):
            meta['name'] = ascii_safe_metadata(item['name'])
        if item.get('age') is not None:
            meta['age'] = str(item['age'])
        if item.get('seoTitle'):
            meta['seoTitle'] = ascii_safe_metadata(item['seoTitle'])
        if item.get('seoDescription'):
            meta['seoDescription'] = ascii_safe_metadata(item['seoDescription'])
        if item.get('city'):
            meta['city'] = ascii_safe_metadata(item['city'])
        if item.get('location'):
            meta['location'] = ascii_safe_metadata(item['location'])
        if item.get('state'):
            meta['state'] = ascii_safe_metadata(item['state'])

        key = update_s3_metadata(s3, str(item['PK']), meta) if meta else None
        if key:
            s3_updated += 1
            print(f'Updated S3 metadata for {item["PK"]} -> {key}')
        else:
            print(f'No matching S3 object for profile {item["PK"]}')

    print(f'DynamoDB rows updated: {rows_updated}')
    print(f'S3 metadata objects updated: {s3_updated}')


if __name__ == '__main__':
    main()
